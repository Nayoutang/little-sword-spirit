from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"E:\GameDev\Projects\Godot\little-sword-spirit")
OUTPUT = ROOT / "documents" / "剑娘_v1.5与当前原型差异说明.docx"

FONT_CN = "Microsoft YaHei"
FONT_EN = "Aptos"
DARK = "263746"
TEAL = "2F7774"
PALE = "EAF4F3"
GRAY = "F3F5F6"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT_EN
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_number(doc, text):
    doc._manual_number_counter = getattr(doc, "_manual_number_counter", 0) + 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.75)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(f"{doc._manual_number_counter}.　{text}")
    set_run_font(run, 9.7)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, DARK)
        set_cell_border(cell)
        set_cell_margins(cell, 120, 130, 120, 130)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(text))
        set_run_font(run, font_size, True, "FFFFFF")
        if widths:
            cell.width = Cm(widths[i])
    for r_index, row_values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            if r_index % 2 == 1:
                set_cell_shading(cell, GRAY)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if i == 0 and len(headers) > 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            set_run_font(run, font_size, i == 0 and len(headers) > 2, "000000")
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_EN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    title.font.name = FONT_EN
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 17, 18, 8),
        ("Heading 2", 13, 13, 6),
        ("Heading 3", 11, 9, 4),
    ):
        style = styles[name]
        style.font.name = FONT_EN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[name]
        style.font.name = FONT_EN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(10.2)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("剑娘项目设计差异说明  |  基线 v1.5  |  2026-09-18")
    set_run_font(run, 8, False, "666666")


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    add_footer(section)
    setup_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("剑娘 v1.5 与当前可玩原型差异说明")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    r = subtitle.add_run("新增内容  设计变更  实现状态  待决策项")
    set_run_font(r, 12, False, TEAL)

    p = doc.add_paragraph()
    p.add_run("文档目的：").bold = True
    p.add_run(
        "以《剑娘 游戏设计策划案 完整版 v1.5》为设计基线，对照 2026 年 9 月 18 日 Godot 4 原型的实际代码与场景，完整记录已经新增、已经改变、仍未实现或仅部分实现的内容，为后续整理 v1.6 策划案提供依据。"
    )
    p = doc.add_paragraph()
    p.add_run("核心结论：").bold = True
    p.add_run(
        "当前原型已经从“核心命题验证稿”发展为具备存档、随机地图、多敌人卡牌战斗、战后奖励、约定、羁绊突破、神通、奇遇与 LLM 对话的完整闭环。与此同时，战斗与养成的结合显著加强，已经突破 v1.5 中“羁绊只解锁内容、绝不提供战斗强度”的原始红线。这是当前版本最重要的设计变化。"
    )

    doc.add_heading("1 对照口径与状态定义", level=1)
    add_table(
        doc,
        ["状态", "含义"],
        [
            ("新增", "v1.5 未提出，当前原型后续增加并已实现。"),
            ("变更", "v1.5 已有方向，但当前规则、数值、流程或定位发生实质调整。"),
            ("细化实现", "v1.5 只有原则或概念，当前已经落实为明确可运行规则。"),
            ("部分实现", "已有原型，但尚未达到 v1.5 描述的完整形态。"),
            ("尚未实现", "v1.5 已规划，但当前工程中没有对应完整系统。"),
        ],
        widths=[3.0, 13.5],
        font_size=9.5,
    )

    doc.add_page_break()
    doc.add_heading("2 总体定位与核心循环变化", level=1)
    add_table(
        doc,
        ["主题", "v1.5 方案", "当前原型", "判断"],
        [
            ("产品中心", "陪伴是主菜，养成是支撑，战斗只是载体。", "整体目标未变，但战斗系统的规则量、游玩时长和奖励频率明显扩大，已经出现抢占陪伴重心的趋势。", "方向保留，重心偏移"),
            ("完整循环", "家中聊天 → 出门叮嘱 → 点状地图 → 战斗 → 战后结算 → 回家。", "存档选择 → Home 自由聊天 → 约定选择 → 10 层地图 → 战斗/问号/宝箱/奇遇 → 战后奖励 → Boss → 结算 → Home。", "大幅细化并闭环"),
            ("一局长度", "目标约 8 至 12 分钟。", "固定单层远征，共 10 个纵向层级；实际时长取决于路线中的战斗数量和多敌人规模。", "结构定型，时长待实测"),
            ("关系与战斗", "关系不直接提升战斗强度。", "羁绊阶段解锁流光与华彩；神通会保命或改变连击规则。", "原则性变更"),
            ("LLM 使用边界", "用于自由聊天与陪伴演绎，避开战斗决策；自由聊天只读不写。", "仍不参与战斗决策；新增奇遇 LLM 演绎与 Boss 后特殊对话语义判定，系统数值继续由本地代码结算。", "边界扩大但仍受控"),
        ],
        widths=[2.4, 4.5, 6.7, 2.4],
        font_size=8.5,
    )

    doc.add_heading("3 当前版本新增内容", level=1)
    doc.add_heading("3.1 三档独立存档", level=2)
    add_bullet(doc, "新增 3 个存档槽，可显示各档羁绊值和关系阶段。")
    add_bullet(doc, "支持删除存档，并将羁绊、关系阶段、神通、特殊事件进度与连续失败次数限制在当前存档内。")
    add_bullet(doc, "支持旧版单存档向第一存档槽迁移。v1.5 未定义多存档结构。")

    doc.add_heading("3.2 一层制 10 层随机地图", level=2)
    add_bullet(doc, "原型曾考虑三层远征，后改为固定一层制，以更快获得通关与关系反馈。")
    add_bullet(doc, "当前地图固定 10 个纵向层级：Home 起点、7 层常规路线、第 9 层固定宝箱、最终 Boss。")
    add_bullet(doc, "地图采用 5 列节点；Home 展开为 5 个起点，层间保留直行并加入单向斜连线，避免 X 形交叉。")
    add_bullet(doc, "同一局使用固定随机种子；进入战斗、事件或奇遇再返回时，可恢复原地图与当前位置。")
    add_bullet(doc, "当前中间 40 个节点配置为：普通战 23、精英 2、宝箱 5、问号 5、奇遇 5。")
    add_bullet(doc, "第一层固定普通战；前三层不出精英；精英在中后段且不连续；第九层固定宝箱。")

    doc.add_heading("3.3 独立奇遇节点与 LLM 演绎", level=2)
    add_bullet(doc, "新增青绿色“奇”节点，与“？”节点完全分离。奇遇不进入战斗，直接进入角色剧情场景。")
    add_bullet(doc, "目前内置雨亭偶憩、无名剑痕、迷路的灯灵 3 个奇遇模板，每个提供 2 个 Galgame 式选择。")
    add_bullet(doc, "选择后由本地代码先锁定并结算羁绊、回血等真实效果，再请求 LLM 生成小墨台词、动作描写和情绪标签。")
    add_bullet(doc, "LLM 无权修改生命、羁绊、卡牌或剧情事实；请求失败、超时或 JSON 格式错误时自动使用本地文案。")
    add_bullet(doc, "Home 与奇遇共用集中式 LLM 配置，API 地址、模型、环境变量和 system prompt 路径不再重复维护。")

    doc.add_heading("3.4 战斗后特殊对话与神通", level=2)
    add_bullet(doc, "新增 3 类 Boss 胜利后特殊经历：残血通关、低生命依靠羁绊技能翻盘、连续失败后通关。")
    add_bullet(doc, "特殊经历回到 Home 后触发 LLM 对话；模型只做语义判断，本地验证 event_id 与 resolved 字段后才解锁神通。")
    add_bullet(doc, "目前神通包括：灵剑护主、剑鸣余韵、不屈剑意，均持久保存在当前存档。")

    doc.add_heading("3.5 数据与工程结构", level=2)
    add_bullet(doc, "脚本已按 autoload、battle、data、home、map、ui 分目录，不再全部堆在 scripts 根目录。")
    add_bullet(doc, "卡牌、敌人、全局平衡、LLM 配置已分别集中到 CardDatabase、EnemyDatabase、BalanceConfig 与 LLMConfig。")
    add_bullet(doc, "运行态由 RunState 管理；神通与特殊对话分别由 AbilityManager 和 SpecialEventManager 管理。")

    doc.add_heading("4 地图与事件系统差异", level=1)
    add_table(
        doc,
        ["项目", "v1.5", "当前原型", "状态"],
        [
            ("地图定位", "仅要求成熟的随机点状路线，地图自身不追求深度。", "路线选择被具体化为 5 列、10 层、分叉与汇流，并加入节点配额和位置约束。", "细化实现"),
            ("节点种类", "普通战、精英战、事件、篝火、Boss。", "普通战、精英、宝箱、问号、奇遇、Boss；没有独立篝火节点。", "变更"),
            ("固定节点", "未规定具体固定楼层；仅提出由浅入深。", "第一层普通战、第九层全宝箱、末层 Boss。", "新增规则"),
            ("问号", "统一作为事件节点概念。", "50% 进入战斗；其中战斗为 70% 普通、30% 精英。其余 50%进入本地事件，可能回血、加上限、选牌、扣 20 HP 或失去一张牌。", "大幅细化"),
            ("宝箱", "未定义独立宝箱节点。", "第九层 5 个节点全部为宝箱；提供回血、增加生命上限并同步回血、三选一卡牌。", "新增"),
            ("奇遇", "没有独立节点，仅有结构化陪伴节点的原则。", "独立“奇”节点，承载角色小剧情、选择、本地效果与 LLM 动态回应。", "新增"),
            ("篝火陪伴", "规划回血、升级、陪她待会儿。", "当前没有篝火和升级牌功能，其关系职责部分被奇遇取代。", "尚未实现并被替代"),
            ("情境皮", "同一功能节点随机套用地名与情境包。", "节点目前仍以战、精、宝、问号、奇等功能字标识，尚无地名和主题情境皮。", "尚未实现"),
        ],
        widths=[2.2, 4.3, 7.0, 2.3],
        font_size=8.5,
    )

    doc.add_heading("5 战斗系统差异", level=1)
    p = doc.add_paragraph()
    p.add_run("整体判断：").bold = True
    p.add_run("战斗从 v1.5 的“够用原型”发展为当前内容最完整的子系统，是新增量最大的部分，也因此成为目前最容易抢走陪伴主题风头的系统。")

    add_table(
        doc,
        ["项目", "v1.5", "当前原型", "状态"],
        [
            ("玩家基础数值", "具体数值待定。", "最大生命 90，精力 3，每回合基础抽 4 张。", "定值"),
            ("牌堆循环", "只说明战斗后三选一构筑。", "完整牌堆、手牌、弃牌堆循环；打出的牌离手，回合结束弃掉剩余手牌，牌堆耗尽后洗弃牌堆；牌堆与弃牌堆可点击查看。", "新增"),
            ("多敌人", "未明确，原型目标偏单一战斗验证。", "普通与精英遭遇均随机 1 至 4 名敌人；攻击牌和指定技能需要选目标；横扫攻击全体。", "新增"),
            ("敌人血量", "未定。", "普通 20 至 40，精英 40 至 60，Boss 200。", "定值"),
            ("敌人意图", "仅提出通过敌人节奏提供喘息回合。", "公开显示下一步：攻击、防御、强化、诅咒或观望；攻击 6 至 10，防御 5 至 9，强化攻击 +2。", "细化实现"),
            ("护盾", "未定。", "玩家与敌人的护盾均在回合结束清空。", "新增规则"),
            ("连击结算", "攻击 +1；防御或状态立即清零；回合末清零；后续攻击递增。", "攻击伤害按出牌前已有连击计算，每层 +2；所有攻击类牌和流光、华彩均吃加成。防御、基础状态清连击；藏锋可令本回合结束不清零。", "变更与扩展"),
            ("特殊技", "连击 2 解锁一次性特殊技。", "特殊技定名“流光”：相识解锁，费 1，基础伤害 15，需 2 连击并消耗 2 连击；常驻按钮，可重复满足条件后使用。", "实质变更"),
            ("终极技", "连击 5 解锁一次性华彩。", "华彩：交心解锁，费 0，基础伤害 35，需 3 连击，使用后连击清零；常驻按钮。", "实质变更"),
            ("胜负流转", "失败重开、Boss 后回家。", "普通胜利进入战后奖励再回地图；Boss 胜利进入奖励再到结算；失败等待 1 秒后进入结算并回 Home。", "细化实现"),
        ],
        widths=[2.1, 4.3, 7.2, 2.2],
        font_size=8.25,
    )

    doc.add_heading("5.1 当前卡牌池", level=2)
    add_table(
        doc,
        ["卡牌", "费用", "当前效果"],
        [
            ("攻击", "1", "基础伤害 6，连击 +1。"),
            ("防御", "1", "格挡 5，清空连击。"),
            ("状态", "1", "抽 1 张，清空连击。"),
            ("强攻", "2", "基础伤害 14，连击 +1。"),
            ("铁壁", "2", "格挡 12，清空连击。"),
            ("横扫", "1", "对所有敌人造成基础伤害 3，连击 +1。"),
            ("叠浪", "1", "基础伤害 6，连击 +2。"),
            ("调息", "0", "抽 1 张，每回合限用一次。"),
            ("掠影", "1", "抽 2 张。"),
            ("破锋", "1", "基础伤害 8，连击 +1，结算后施加 2 层易伤；每层使受到伤害 +1。"),
            ("卸力", "1", "指定敌人，使其下一次攻击伤害降低 5。"),
            ("藏锋", "1", "格挡 6，本回合结束时连击不清零。"),
            ("诅咒", "0", "无法打出，由敌人诅咒意图加入弃牌堆。"),
            ("流光", "1", "基础伤害 15，需 2 连击，消耗 2 连击，相识解锁。"),
            ("华彩", "0", "基础伤害 35，需 3 连击，连击清零，交心解锁。"),
        ],
        widths=[3.0, 2.0, 11.2],
        font_size=9.0,
    )
    p = doc.add_paragraph()
    p.add_run("初始牌组：").bold = True
    p.add_run("攻击 5、防御 5、状态 1、铁壁 1、强攻 1、横扫 1、叠浪 1，共 15 张。流光和华彩作为常驻羁绊技能按钮，不进入普通抽牌循环。")

    doc.add_heading("6 羁绊 约定 与成长系统差异", level=1)
    add_table(
        doc,
        ["项目", "v1.5", "当前原型", "状态"],
        [
            ("羁绊阶段", "0 至 100，提议 L0 至 L5 或若干里程碑。", "固定 4 阶段：初遇 0、相识 25、交心 50、生死之交 75。", "具体化"),
            ("关系突破", "达到临界后通过懂不懂她的选择跨阶；选错暂缓，不扣永久羁绊。", "相识、交心、生死之交各有三选一突破场景；选错暂缓并保留羁绊。", "已实现主体"),
            ("短期壳变硬", "选错或被冷落会产生临时态度回冷。", "尚无独立短期情绪或壳硬度变量；只有关系阶段提示词。", "尚未实现"),
            ("羁绊增长", "主要来自约定兑现、偶发陪伴选择和里程碑；自由聊天不增加。", "每场战斗固定 +2；protect 约定每层满足可 +6；finish 通关可 +10；奇遇选择 +1 或 +2。自由聊天仍不直接增加。", "明显扩大"),
            ("约定种类", "Demo 计划 1 至 2 种，如 protect、finish。", "已实现 protect 与 finish。protect 以生命不低于最大生命 25%为判定，按战斗层结算；finish 在整局结算。", "已实现并细化"),
            ("羁绊回报", "只解锁内容与关系表现，不直接提高战斗强度。", "除关系台词和突破外，还解锁流光、华彩；特殊对话可解锁改变战斗规则的神通。", "核心原则改变"),
            ("战后奖励", "战后三选一；偶发可放弃强化换取陪伴和羁绊。", "当前每场战斗后在回血 10 或三选一卡牌之间选择；羁绊 +2 自动发生，没有“陪她待一会儿”的机会成本选项。", "变更"),
        ],
        widths=[2.2, 4.4, 7.0, 2.2],
        font_size=8.35,
    )

    doc.add_heading("7 LLM 与记忆系统差异", level=1)
    add_table(
        doc,
        ["项目", "v1.5", "当前原型", "状态"],
        [
            ("Home 自由聊天", "玩家自由输入，LLM 读取人设、态度、记忆，只读不写。", "已接入兼容 OpenAI Chat Completions 的 API；读取小墨 system prompt、关系阶段、已掌握神通与当前特殊事件。", "已实现主体"),
            ("长期记忆", "保存关键事实，如上次约定、冷落次数、共同击败的敌人、羁绊阶段。", "当前持久保存羁绊、阶段、神通、特殊事件和连续失败次数；普通聊天 messages 只存在于当前 Home 场景会话，未形成通用关键事实清单。", "部分实现"),
            ("结构化写入", "关键节点使用预设选项，本地写入系统。", "约定、战后奖励、关系突破、事件和奇遇均由本地代码写入；原则保留。", "已实现"),
            ("LLM 特殊事件", "未定义。", "Boss 战表现会排队形成特殊对话，LLM 判断玩家回应是否满足语义条件，通过本地校验后解锁神通。", "新增"),
            ("LLM 奇遇", "情境节点可由 LLM 生成短模板，但未定义执行方案。", "奇遇背景与选项固定，本地先结算，LLM 仅生成后续台词与动作；失败自动回退本地文案。", "新增并受控"),
            ("集中配置", "未涉及工程配置。", "API 地址、模型、密钥环境变量、试玩密钥与 system prompt 路径统一由 LLMConfig 管理。", "新增"),
        ],
        widths=[2.2, 4.4, 7.0, 2.2],
        font_size=8.4,
    )

    doc.add_heading("8 v1.5 已规划但当前仍缺失的内容", level=1)
    doc._manual_number_counter = 0
    add_number(doc, "完整关键事实记忆系统。当前没有把玩家名字、偏好、过去自由对话中的重要事实提炼并跨场景、跨启动保存。")
    add_number(doc, "独立的壳硬度与短期情绪修正。关系阶段会改变 prompt，但被冷落、刚守约或突破失败造成的临时态度变化尚未建模。")
    add_number(doc, "4 至 6 个表情差分及随态度驱动的表现。当前主要依赖固定画面与文字，没有完整表情状态机。")
    add_number(doc, "战斗 Q 版动作、濒危写实 cut-in 与高光反差演出。战斗目前以文字和色块验证机制。")
    add_number(doc, "篝火节点、卡牌升级与“陪她待会儿”的局内陪伴选择。当前由宝箱、问号与奇遇承担部分节奏功能。")
    add_number(doc, "地图情境皮与地名包。节点功能已经成熟，但尚未包装为后山、镇上、禁地等主题化随机情境。")
    add_number(doc, "“冷落”行为及其长期记录。当前没有 neglect_count 或等价系统。")
    add_number(doc, "多剑娘框架仍未实现，这与 v1.5 Demo 明确排除项一致。")

    doc.add_heading("9 与 v1.5 核心原则不一致的关键决策", level=1)
    doc.add_heading("9.1 羁绊已开始提供战斗能力", level=2)
    p = doc.add_paragraph()
    p.add_run("v1.5 原则：").bold = True
    p.add_run("羁绊只购买“她”和关系内容，不购买力量，避免跨局积累削弱 Roguelike 每局构筑的张力。")
    p = doc.add_paragraph()
    p.add_run("当前实现：").bold = True
    p.add_run("相识解锁流光，交心解锁华彩；灵剑护主、剑鸣余韵、不屈剑意均会直接改变战斗结果。")
    p = doc.add_paragraph()
    p.add_run("影响：").bold = True
    p.add_run("关系成长获得了更直观的玩法反馈，但也使老存档天然强于新存档，并使战斗成为刷羁绊和验证成长的主要场所。该变化应在 v1.6 中正式确认为新方向，或重新收回为表现与内容型奖励。")

    doc.add_heading("9.2 战斗羁绊增长过于高频", level=2)
    p = doc.add_paragraph(
        "v1.5 主张每局只有少量高质量陪伴写入点；当前每场战斗固定增加 2 羁绊，protect 约定还能按层反复结算。结果是玩家提高羁绊的最有效方式可能变成多打战斗，而不是理解小墨或兑现有重量的约定。"
    )

    doc.add_heading("9.3 结构化陪伴频率已经提高", level=2)
    p = doc.add_paragraph(
        "v1.5 建议每局约 2 至 3 次结构化陪伴交互；当前地图配置 5 个奇遇节点，玩家虽不会全部经过，但单局可能多次进入奇遇。奇遇能够削弱战斗占比并突出角色，但应通过路线概率、事件冷却或不重复池控制密度。"
    )

    doc.add_heading("10 建议写入 v1.6 的正式结论", level=1)
    doc._manual_number_counter = 0
    add_number(doc, "正式确认 Demo 为单层 10 层短局结构，不再保留三层爬塔设想。")
    add_number(doc, "将节点类型更新为普通战、精英、宝箱、问号、奇遇、Boss，并明确篝火是否永久取消。")
    add_number(doc, "将奇遇确立为陪伴系统的局内主入口，明确固定选项、本地结算、LLM 演绎、失败回退的技术边界。")
    add_number(doc, "重新裁决“羁绊能否提供战力”。若保留，应改写 v1.5 红线，并限制增益为横向玩法或小幅协同，避免形成数值碾压。")
    add_number(doc, "降低“每场战斗固定加羁绊”的权重，把主要羁绊收益迁回约定、奇遇、关系突破和 Boss 后特殊对话。")
    add_number(doc, "把持久记忆和短期态度系统列为下一阶段最高优先级；它们比继续扩充卡牌更直接服务核心命题。")
    add_number(doc, "为地图增加奇遇去重和事件冷却，避免同一局反复抽到相同三组剧情。")
    add_number(doc, "补充试玩指标：单局实际时长、平均战斗节点数、奇遇触发数、羁绊增长速度、玩家是否记得小墨的反应而非卡牌数值。")

    doc.add_page_break()
    doc.add_heading("附录 当前原型场景与系统清单", level=1)
    add_table(
        doc,
        ["模块", "当前内容"],
        [
            ("存档入口", "save_select.tscn：3 个独立存档槽与删除确认。"),
            ("Home", "home.tscn：自由 LLM 对话、羁绊显示、神通与特殊事件上下文、出击与存档入口。"),
            ("约定", "promise.tscn：protect、finish 或不立约定。"),
            ("地图", "map.tscn：10 层、5 列、随机节点类型、路线分叉、进度恢复。"),
            ("战斗", "battle.tscn：多敌人、选目标、牌堆循环、连击、意图、状态、流光与华彩。"),
            ("战后奖励", "battle_reward.tscn：回血或三选一卡牌，之后返回地图或结算。"),
            ("宝箱与问号", "event.tscn：回血、加生命上限、选牌及问号负面结果。"),
            ("奇遇", "adventure.tscn：固定选择、本地效果、LLM 动态演绎与离线兜底。"),
            ("关系突破", "relationship_milestone.tscn：三个阶段的理解选择。"),
            ("远征结算", "settlement.tscn：胜负、约定结果、羁绊状态及特殊对话提示。"),
        ],
        widths=[4.0, 12.2],
        font_size=9.0,
    )

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.name is None:
                set_run_font(run)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
