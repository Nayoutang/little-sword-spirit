from pathlib import Path
from docx import Document


source = Path(r"E:\GameDev\Documents\GameDesign\剑娘_游戏设计策划案_完整版_v1.5.docx")
document = Document(source)

print("PARAGRAPHS")
for index, paragraph in enumerate(document.paragraphs):
    if not paragraph.text.strip():
        continue
    style_name = paragraph.style.name if paragraph.style is not None else "None"
    print(f"{index:04d}\t{style_name}\t{paragraph.text}")

print("TABLES")
for table_index, table in enumerate(document.tables):
    print(f"TABLE {table_index}")
    for row in table.rows:
        print(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
